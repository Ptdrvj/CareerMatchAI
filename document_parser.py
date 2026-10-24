"""
Multi-format Document Ingestion Pipeline
Parses .pdf, .docx, and .txt files to extract unstructured text with automated fallback handling.
"""

import io
import os
import re
from typing import Dict, Any, Union, Optional


class DocumentParserError(Exception):
    """Custom exception raised when document parsing fails."""
    pass


class DocumentParser:
    """
    Robust document ingestion pipeline capable of parsing PDF, DOCX, and TXT files.
    Includes encoding fallbacks, structural table parsing for DOCX, and text normalization.
    """

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

    @classmethod
    def clean_text(cls, text: str) -> str:
        """
        Normalizes extracted text by removing control characters,
        collapsing multiple blank lines, and stripping trailing whitespace.
        """
        if not text:
            return ""
        # Remove null characters and non-printable control characters (keep newlines/tabs)
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
        # Normalize carriage returns
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        # Collapse 3+ consecutive newlines into 2
        text = re.sub(r"\n{3,}", "\n\n", text)
        # Collapse multiple inline spaces
        text = re.sub(r"[ \t]{2,}", " ", text)
        return text.strip()

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> str:
        """
        Parses PDF using PyPDF with error recovery and page-by-page extraction.
        """
        try:
            import pypdf
        except ImportError:
            raise DocumentParserError("pypdf is required to parse PDF documents. Install with 'pip install pypdf'.")

        stream = io.BytesIO(file_bytes)
        try:
            reader = pypdf.PdfReader(stream)
            if reader.is_encrypted:
                try:
                    # Attempt empty password decryption
                    reader.decrypt("")
                except Exception:
                    raise DocumentParserError("The uploaded PDF is password-protected. Please provide an unencrypted PDF.")

            extracted_pages = []
            for i, page in enumerate(reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        extracted_pages.append(page_text)
                except Exception as page_err:
                    # Continue extracting other pages even if one fails
                    continue

            full_text = "\n\n".join(extracted_pages)
            cleaned = cls.clean_text(full_text)

            if not cleaned:
                raise DocumentParserError(
                    "The PDF appears to be empty or contains only scanned images without an OCR text layer. "
                    "Please upload a searchable text-based PDF or copy/paste the text directly."
                )
            return cleaned

        except DocumentParserError:
            raise
        except Exception as e:
            raise DocumentParserError(f"Failed to read PDF document: {str(e)}")

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> str:
        """
        Parses DOCX documents extracting both body paragraphs and table contents.
        Many resumes format skills, employment history, and education in tables.
        """
        try:
            import docx
        except ImportError:
            raise DocumentParserError("python-docx is required to parse DOCX documents. Install with 'pip install python-docx'.")

        stream = io.BytesIO(file_bytes)
        try:
            doc = docx.Document(stream)
            content_blocks = []

            # 1. Extract paragraphs
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    content_blocks.append(text)

            # 2. Extract tables (vital for resumes formatted in columnar layout)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    # Deduplicate adjacent cells if merged
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    if unique_cells:
                        content_blocks.append(" | ".join(unique_cells))

            full_text = "\n".join(content_blocks)
            cleaned = cls.clean_text(full_text)

            if not cleaned:
                raise DocumentParserError("The DOCX document contains no readable text.")
            return cleaned

        except DocumentParserError:
            raise
        except Exception as e:
            raise DocumentParserError(f"Failed to read DOCX document: {str(e)}")

    @classmethod
    def parse_txt(cls, file_bytes: bytes) -> str:
        """
        Parses plain text / markdown files with automated fallback across common encodings.
        """
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
        for enc in encodings:
            try:
                decoded = file_bytes.decode(enc)
                cleaned = cls.clean_text(decoded)
                if cleaned:
                    return cleaned
            except UnicodeDecodeError:
                continue

        raise DocumentParserError("Could not decode text file with standard encodings (UTF-8, Latin-1, CP1252).")

    @classmethod
    def extract_text(cls, file_obj: Union[io.BytesIO, bytes, Any], filename: str = "") -> Dict[str, Any]:
        """
        Unified extraction method handling file-like objects (e.g. Streamlit UploadedFile)
        or raw bytes. Returns extracted text and document statistics.
        """
        if hasattr(file_obj, "read"):
            file_bytes = file_obj.read()
            if hasattr(file_obj, "seek"):
                file_obj.seek(0)
            if not filename and hasattr(file_obj, "name"):
                filename = file_obj.name
        elif isinstance(file_obj, bytes):
            file_bytes = file_obj
        else:
            raise DocumentParserError(f"Unsupported file input type: {type(file_obj)}")

        if not file_bytes:
            raise DocumentParserError("Uploaded file is empty (0 bytes).")

        ext = os.path.splitext(filename.lower())[1] if filename else ""

        extracted_text = ""
        # Route by extension, with automatic fallback testing
        if ext == ".pdf":
            extracted_text = cls.parse_pdf(file_bytes)
        elif ext == ".docx":
            extracted_text = cls.parse_docx(file_bytes)
        elif ext in [".txt", ".md"]:
            extracted_text = cls.parse_txt(file_bytes)
        else:
            # Fallback heuristic: Try PDF magic bytes (%PDF)
            if file_bytes.startswith(b"%PDF"):
                extracted_text = cls.parse_pdf(file_bytes)
            # Try DOCX magic bytes (PK zip header)
            elif file_bytes.startswith(b"PK\x03\x04"):
                try:
                    extracted_text = cls.parse_docx(file_bytes)
                except Exception:
                    extracted_text = cls.parse_txt(file_bytes)
            else:
                # Default to text decoding
                extracted_text = cls.parse_txt(file_bytes)

        word_count = len(extracted_text.split())
        char_count = len(extracted_text)

        return {
            "filename": filename or "uploaded_document",
            "extension": ext or "unknown",
            "text": extracted_text,
            "word_count": word_count,
            "char_count": char_count,
            "byte_size": len(file_bytes),
        }
